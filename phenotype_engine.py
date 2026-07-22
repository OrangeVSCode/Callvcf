#!/usr/bin/env python3
"""Phenotype QC, BLUE/BLUP estimates and visual reports for GPA-Accelerator."""

import base64
import csv
import html
import json
import math
import os
import re
import statistics
import time
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path

from vcf_service import VCFError


MISSING = {"", ".", "na", "nan", "null", "none", "n/a", "-"}
ALIASES = {
    "sample": ["sample", "sample_id", "material", "material_id", "accession", "genotype", "line", "id", "材料", "材料名", "品种", "样本"],
    "trait": ["trait", "phenotype", "variable", "指标", "性状", "表型"],
    "value": ["value", "measurement", "phenotype_value", "数值", "值", "表型值"],
    "year": ["year", "season", "年份", "年度"],
    "location": ["location", "site", "environment", "env", "地点", "试点", "环境"],
    "latitude": ["latitude", "lat", "纬度"],
    "longitude": ["longitude", "lon", "lng", "经度"],
    "replicate": ["replicate", "rep", "block", "重复", "区组"],
    "group": ["group", "population", "period4", "群体", "分组"],
}


# name, unit, single-plant range, population-mean range, confidence.
# Undefined traits from the guide deliberately have no invented biological limits.
TRAIT_INFO = {
    "PH": ("株高", "cm", (60, 200), (80, 100), "medium"),
    "FBH": ("第一果枝高度", "cm", (10, 60), (20, 60), "medium"),
    "FFN": ("第一果枝节位", "节", (4, 7), (5, 6), "medium"),
    "FBN": ("果枝数", "个/株", (4, 15), (8, 12), "medium"),
    "FNN": ("果节数", "个/株", (10, 60), (10, 50), "low"),
    "BN": ("结铃数", "个/株", (5, 15), (6, 12), "medium"),
    "SBW": ("单铃重", "g", (3, 8), (5, 7), "medium"),
    "LP": ("衣分", "%", (25, 45), (38, 42), "medium"),
    "SPAD": ("SPAD值", "", (30, 60), (35, 50), "low"),
    "LA": ("叶面积", "cm²", (50, 200), None, "low"),
    "FL": ("纤维长度", "mm", (20, 40), (26, 32), "medium"),
    "UNI": ("纤维整齐度", "%", (75, 90), (83, 86), "medium"),
    "FS": ("纤维强力", "cN/tex", (25, 40), (28, 32), "medium"),
    "MIC": ("马克隆值", "", (3.5, 5.5), (4.0, 4.5), "medium"),
    "EL": ("纤维伸长率", "%", (4, 8), (5, 6), "medium"),
    "LN": ("籽指", "g", (8, 15), (10, 12), "medium"),
    "FBA": ("果枝夹角", "°", (15, 60), None, "low"),
    "LIA": ("衣指", "g", (5, 12), (7, 9), "medium"),
    "PD": ("生育期", "天/节点", None, None, "undefined"),
    "HBR": ("HBR（定义待确认）", "", None, None, "undefined"),
    "NLB": ("NLB（定义待确认）", "", None, None, "undefined"),
    "LR": ("LR（定义待确认）", "", None, None, "undefined"),
    "NDLHB": ("NDLHB（定义待确认）", "", None, None, "undefined"),
    "TNB": ("TNB（可能为总铃数）", "", None, None, "undefined"),
    "LC": ("LC（定义待确认）", "", None, None, "undefined"),
    "SCW": ("SCW（定义待确认）", "", None, None, "undefined"),
    "FBL": ("FBL（定义待确认）", "", None, None, "undefined"),
    "LL": ("叶片长度/LL", "cm", None, None, "undefined"),
    "LW": ("叶片宽度/LW", "cm", None, None, "undefined"),
    "LAL": ("LAL（定义待确认）", "cm", None, None, "undefined"),
    "LAW": ("LAW（定义待确认）", "cm", None, None, "undefined"),
}


COTTON_SPECIES = {
    "cotton_upland": {
        "name": "陆地棉", "scientific_name": "Gossypium hirsutum", "ploidy": 4,
        "notes": "AD1异源四倍体；PH和FL使用指南中的陆地棉细化范围。",
        "overrides": {"PH": ((60, 160), (80, 100), "medium"), "FL": ((20, 35), (26, 32), "medium")},
    },
    "cotton_barbadense": {
        "name": "海岛棉", "scientific_name": "Gossypium barbadense", "ploidy": 4,
        "notes": "AD2异源四倍体；长绒棉株高与纤维长度使用更宽范围。",
        "overrides": {"PH": ((80, 200), (100, 180), "low"), "FL": ((30, 40), (30, 40), "medium")},
    },
    "cotton_herbaceum": {
        "name": "草棉", "scientific_name": "Gossypium herbaceum", "ploidy": 2,
        "notes": "A1二倍体；多数性状缺少草棉专属范围，使用跨棉种保守范围并降低置信度。",
        "overrides": {},
    },
}


def _norm(value):
    return re.sub(r"[\s_\-./()（）]+", "", str(value or "").strip().lower())


def _number(value):
    text = str(value if value is not None else "").strip()
    if text.lower() in MISSING:
        return None
    text = text.replace(",", "")
    if text.endswith("%"):
        text = text[:-1]
    try:
        value = float(text)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def _quantile(values, q):
    data = sorted(values)
    if not data:
        return None
    if len(data) == 1:
        return data[0]
    index = (len(data) - 1) * q
    lo, hi = int(math.floor(index)), int(math.ceil(index))
    return data[lo] if lo == hi else data[lo] * (hi - index) + data[hi] * (index - lo)


def _safe_sd(values):
    return statistics.stdev(values) if len(values) > 1 else 0.0


def _describe(values):
    if not values:
        return {"n": 0}
    n = len(values); mean = statistics.fmean(values); median = statistics.median(values); sd = _safe_sd(values)
    q1, q3 = _quantile(values, .25), _quantile(values, .75)
    mad = statistics.median([abs(x - median) for x in values])
    trim = int(n * .05); ordered = sorted(values)
    trimmed = ordered[trim:n-trim] if trim and n-trim > trim else ordered
    skew = sum(((x - mean) / sd) ** 3 for x in values) * n / ((n - 1) * (n - 2)) if sd and n > 2 else 0.0
    kurtosis = 0.0
    if sd and n > 3:
        raw = sum(((x - mean) / sd) ** 4 for x in values)
        kurtosis = n * (n + 1) * raw / ((n - 1) * (n - 2) * (n - 3)) - 3 * (n - 1) ** 2 / ((n - 2) * (n - 3))
    return {
        "n": n, "mean": mean, "median": median, "trimmed_mean_5pct": statistics.fmean(trimmed),
        "sd": sd, "variance": sd * sd, "se": sd / math.sqrt(n), "cv": sd / abs(mean) if mean else None,
        "mad": mad, "iqr": q3 - q1, "min": min(values), "q05": _quantile(values, .05),
        "q25": q1, "q75": q3, "q95": _quantile(values, .95), "max": max(values),
        "skewness": skew, "excess_kurtosis": kurtosis,
    }


def _read_table(path, sheet=None):
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise VCFError("读取XLSX需要openpyxl；请安装后重试，或另存为CSV/TSV")
        book = load_workbook(path, read_only=True, data_only=True)
        name = sheet if sheet in book.sheetnames else book.sheetnames[0]
        iterator = book[name].iter_rows(values_only=True)
        try:
            headers = [str(x or "").strip() for x in next(iterator)]
        except StopIteration:
            raise VCFError("表型文件为空")
        rows = [dict(zip(headers, values)) for values in iterator if any(x is not None and str(x).strip() for x in values)]
        sheet_names = list(book.sheetnames)
        book.close()
        return headers, rows, {"format": "xlsx", "sheet": name, "sheets": sheet_names}
    if suffix == ".xls":
        raise VCFError("旧版.XLS请另存为.XLSX、CSV或TSV")
    if suffix not in {".csv", ".tsv", ".txt"}:
        raise VCFError("表型文件支持 .xlsx、.csv、.tsv 和 .txt")
    raw = path.read_bytes(); text = None; encoding = None
    for candidate in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(candidate); encoding = candidate; break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise VCFError("无法识别表型文件编码；建议另存为UTF-8 CSV/TSV")
    delimiter = "\t" if suffix == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",\t;").delimiter
    except csv.Error:
        pass
    reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
    headers = [str(x or "").strip() for x in (reader.fieldnames or [])]
    rows = [{str(k or "").strip(): v for k, v in row.items()} for row in reader if any(str(x or "").strip() for x in row.values())]
    return headers, rows, {"format": suffix.lstrip("."), "encoding": encoding, "delimiter": delimiter}


def _detect_column(headers, role, explicit=None):
    if explicit and explicit in headers:
        return explicit
    lookup = {_norm(x): x for x in headers}
    return next((lookup[_norm(alias)] for alias in ALIASES[role] if _norm(alias) in lookup), None)


def _parse_custom_thresholds(text):
    output = {}
    for line_no, line in enumerate(str(text or "").splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        bits = [x.strip() for x in re.split(r"[,\t:]", line)]
        if len(bits) not in {3, 5}:
            raise VCFError("自定义阈值第{}行应为 trait,single_min,single_max[,mean_min,mean_max]".format(line_no))
        nums = [_number(x) for x in bits[1:]]
        if any(x is None for x in nums):
            raise VCFError("自定义阈值第{}行包含非数值范围".format(line_no))
        output[bits[0].upper()] = {"single": nums[:2], "mean": nums[2:] or None, "confidence": "user"}
    return output


def _threshold_for(trait, species_id, custom):
    key = trait.upper()
    if key in custom:
        return {"trait": key, "source": "用户自定义", **custom[key]}
    info = TRAIT_INFO.get(key)
    if not info:
        return {"trait": key, "single": None, "mean": None, "confidence": "unknown", "source": "无内置定义"}
    _, _, single, mean, confidence = info
    species = COTTON_SPECIES.get(species_id)
    source = "棉花表型阈值规划（跨棉种保守范围）"
    if species and key in species["overrides"]:
        single, mean, confidence = species["overrides"][key]
        source = "棉花表型阈值规划（{}物种细化）".format(species["name"])
    elif species_id == "cotton_herbaceum" and confidence not in {"undefined", "unknown"}:
        confidence = "low"; source += "；草棉专属证据不足"
    return {"trait": key, "single": list(single) if single else None, "mean": list(mean) if mean else None, "confidence": confidence, "source": source}


def phenotype_catalog():
    traits = []
    for code, (name, unit, single, mean, confidence) in TRAIT_INFO.items():
        traits.append({"code": code, "name": name, "unit": unit, "single_range": single, "mean_range": mean, "confidence": confidence})
    return {
        "cotton_species": [{"id": key, **{k: v for k, v in value.items() if k != "overrides"}} for key, value in COTTON_SPECIES.items()],
        "cotton_traits": traits, "supported_formats": [".xlsx", ".csv", ".tsv", ".txt"],
        "outlier_methods": ["IQR", "3-sigma", "MAD robust-z", "tail percentile", "biological range"],
        "blup_method": "environment-centered empirical random-entry BLUP; variance components exported",
    }


def _environment(row):
    value = " × ".join(x for x in (str(row.get("year") or "").strip(), str(row.get("location") or "").strip()) if x)
    return value or "整体"


def _blue_blup(rows):
    env_values = defaultdict(list)
    for row in rows:
        env_values[_environment(row)].append(row["value"])
    grand = statistics.fmean([x["value"] for x in rows]); env_mean = {k: statistics.fmean(v) for k, v in env_values.items()}
    adjusted, raw = defaultdict(list), defaultdict(list)
    for row in rows:
        adjusted[row["sample"]].append(row["value"] - env_mean[_environment(row)] + grand)
        raw[row["sample"]].append(row["value"])
    means = {key: statistics.fmean(values) for key, values in adjusted.items()}
    total, groups = sum(map(len, adjusted.values())), len(adjusted)
    residual = sum(sum((x-means[key])**2 for x in values) for key, values in adjusted.items()) / max(1, total-groups)
    overall = sum(len(adjusted[k])*means[k] for k in adjusted) / max(1, total)
    ms_between = sum(len(v)*(means[k]-overall)**2 for k, v in adjusted.items()) / max(1, groups-1)
    n0 = (total - sum(len(v)**2 for v in adjusted.values())/max(1,total)) / max(1,groups-1)
    genetic = max(0.0, (ms_between-residual)/n0) if n0 else 0.0
    harmonic = groups / sum(1/len(v) for v in adjusted.values()) if groups else 1
    h2 = genetic/(genetic+residual/harmonic) if genetic+residual/harmonic else 0.0
    estimates = []
    for sample in sorted(adjusted):
        n = len(adjusted[sample]); reliability = n*genetic/(n*genetic+residual) if n*genetic+residual else 0.0
        blue = means[sample]
        estimates.append({"sample": sample, "n": n, "raw_mean": statistics.fmean(raw[sample]), "blue": blue, "blup": overall+reliability*(blue-overall), "blup_reliability": reliability})
    variance = {
        "genetic_variance": genetic, "residual_variance": residual, "broad_sense_h2_entry_mean": h2,
        "harmonic_replicates": harmonic, "environments": len(env_mean),
        "method": "环境固定效应中心化后的BLUE；随机材料效应经验BLUP（方法矩估计方差组分）",
        "caveat": "用于快速质控与排序；复杂空间设计、亲缘矩阵或非平衡G×E研究应使用REML模型复核。",
    }
    return estimates, variance


def _linear_trend(points):
    numeric = []
    for key, value in points:
        match = re.search(r"-?\d+(?:\.\d+)?", str(key))
        if match: numeric.append((float(match.group()), value))
    if len(numeric) < 3: return None
    xs, ys = zip(*numeric); xm, ym = statistics.fmean(xs), statistics.fmean(ys); den = sum((x-xm)**2 for x in xs)
    return sum((x-xm)*(y-ym) for x,y in numeric)/den if den else 0.0


def _outlier_rows(rows, config, threshold):
    values = [x["value"] for x in rows]
    if len(values) < 4: return [], {"reason": "少于4个观测值"}
    mean, sd, median = statistics.fmean(values), _safe_sd(values), statistics.median(values)
    q1, q3 = _quantile(values,.25), _quantile(values,.75); factor = float(config.get("iqr_factor",1.5))
    iqr_bounds = (q1-factor*(q3-q1), q3+factor*(q3-q1)); sigma = float(config.get("sigma",3.0))
    sigma_bounds = (mean-sigma*sd, mean+sigma*sd); mad = statistics.median([abs(x-median) for x in values]); mad_z = float(config.get("mad_z",3.5))
    tail = min(.10,max(0.0,float(config.get("tail_fraction",.005)))); tail_bounds = (_quantile(values,tail),_quantile(values,1-tail)) if tail else (None,None)
    required = max(1,int(config.get("consensus",2))); output=[]
    for index,row in enumerate(rows,1):
        value=row["value"]; methods=[]
        if value<iqr_bounds[0] or value>iqr_bounds[1]: methods.append("IQR")
        if sd and (value<sigma_bounds[0] or value>sigma_bounds[1]): methods.append("3-sigma")
        if mad and abs(.67448975*(value-median)/mad)>mad_z: methods.append("MAD robust-z")
        if tail and len(values)>=40 and (value<tail_bounds[0] or value>tail_bounds[1]): methods.append("tail percentile")
        biological=threshold.get("single")
        if biological and (value<biological[0] or value>biological[1]): methods.append("biological range")
        statistical=sum(x!="biological range" for x in methods)
        biological_standalone = threshold.get("confidence") in {"medium", "high", "user"}
        if methods and (statistical>=required or ("biological range" in methods and biological_standalone)):
            output.append({"row":index,"sample":row["sample"],"trait":row["trait"],"value":value,"year":row.get("year") or "","location":row.get("location") or "","methods":";".join(methods),"method_count":len(methods),"consensus_statistical":statistical>=required,"severity":"critical" if statistical>=max(3,required+1) else "warning"})
    return output,{"iqr_bounds":list(iqr_bounds),"sigma_bounds":list(sigma_bounds),"mad_center":median,"mad":mad,"mad_z_threshold":mad_z,"tail_bounds":list(tail_bounds),"consensus_required":required}


def _summary_groups(rows, field):
    grouped=defaultdict(list)
    for row in rows:
        if str(row.get(field) or "").strip(): grouped[str(row[field])].append(row["value"])
    return [{field:key,**_describe(values)} for key,values in sorted(grouped.items())]


def _slug(text):
    return (re.sub(r"[^A-Za-z0-9_.-]+","_",str(text)).strip("._")[:70] or "trait")


def _pearson(xs,ys):
    if len(xs)<3: return None
    xm,ym=statistics.fmean(xs),statistics.fmean(ys); denx=sum((x-xm)**2 for x in xs); deny=sum((y-ym)**2 for y in ys)
    return sum((x-xm)*(y-ym) for x,y in zip(xs,ys))/math.sqrt(denx*deny) if denx and deny else None


def _ranks(values):
    ordered=sorted(range(len(values)),key=lambda i:values[i]); ranks=[0.0]*len(values); index=0
    while index<len(ordered):
        end=index+1
        while end<len(ordered) and values[ordered[end]]==values[ordered[index]]: end+=1
        rank=(index+1+end)/2
        for pos in ordered[index:end]: ranks[pos]=rank
        index=end
    return ranks


def _trait_correlations(observations):
    grouped=defaultdict(list)
    for row in observations: grouped[(row["sample"],row["trait"])].append(row["value"])
    matrix=defaultdict(dict)
    for (sample,trait),values in grouped.items(): matrix[trait][sample]=statistics.fmean(values)
    traits=sorted(matrix); output=[]
    for i,left in enumerate(traits):
        for right in traits[i:]:
            common=sorted(set(matrix[left])&set(matrix[right])); xs=[matrix[left][x] for x in common]; ys=[matrix[right][x] for x in common]
            output.append({"trait_x":left,"trait_y":right,"n":len(common),"pearson_r":_pearson(xs,ys),"spearman_rho":_pearson(_ranks(xs),_ranks(ys)) if len(common)>=3 else None})
    return output


def _svg_correlations(rows):
    traits=sorted({x["trait_x"] for x in rows}|{x["trait_y"] for x in rows}); n=len(traits)
    if n<2: return "<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 700 220'><text x='30' y='50'>至少需要两个表型才能计算相关性。</text></svg>"
    cell=max(18,min(46,620/max(1,n))); size=cell*n; width=size+150; height=size+130; lookup={}
    for row in rows: lookup[(row["trait_x"],row["trait_y"])]=lookup[(row["trait_y"],row["trait_x"])]=row["pearson_r"]
    parts=["<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 {0:.0f} {1:.0f}'><rect width='100%' height='100%' fill='white'/><text x='85' y='28' font-size='18' font-weight='700' fill='#173b31'>表型 Pearson 相关热图</text>".format(width,height)]
    for i,trait in enumerate(traits):
        x=85+i*cell+cell/2; y=60+i*cell+cell/2
        parts.append("<text x='{:.1f}' y='52' transform='rotate(-55 {:.1f} 52)' text-anchor='start' font-size='9' fill='#566b62'>{}</text>".format(x,x,html.escape(trait[:16])))
        parts.append("<text x='78' y='{:.1f}' text-anchor='end' font-size='9' fill='#566b62'>{}</text>".format(y+3,html.escape(trait[:16])))
        for j,other in enumerate(traits):
            value=lookup.get((trait,other)); q=0 if value is None else value; red=int(245-135*max(0,q)); blue=int(245-125*max(0,-q)); green=int(245-95*abs(q)); color="rgb({},{},{})".format(red,green,blue)
            parts.append("<rect x='{:.1f}' y='{:.1f}' width='{:.1f}' height='{:.1f}' fill='{}' stroke='white'><title>{} × {}: {}</title></rect>".format(85+j*cell,60+i*cell,cell,cell,color,html.escape(trait),html.escape(other),"NA" if value is None else "{:.3f}".format(value)))
    parts.append("</svg>"); return "".join(parts)


def _svg_density(series, title, unit=""):
    palette = ["#176b5a", "#d07a24", "#73559b", "#3f79b5", "#b44e5a", "#67964e", "#2d8b91"]
    all_values = [x for _, values in series for x in values]
    if not all_values:
        return "<svg xmlns='http://www.w3.org/2000/svg' width='900' height='430'><text x='40' y='60'>无可绘制数据</text></svg>"
    lo, hi = min(all_values), max(all_values)
    if lo == hi: lo, hi = lo-.5, hi+.5
    pad = (hi-lo)*.08; lo, hi = lo-pad, hi+pad
    xs = [lo+(hi-lo)*i/119 for i in range(120)]; curves=[]; max_y=0
    for label, values in series:
        sampled = values if len(values)<=5000 else values[::max(1,len(values)//5000)]
        sd=_safe_sd(sampled); bandwidth=1.06*sd*(len(sampled)**-.2) if sd else (hi-lo)/30
        bandwidth=max(bandwidth,(hi-lo)/250); scale=1/(len(sampled)*bandwidth*math.sqrt(2*math.pi))
        ys=[scale*sum(math.exp(-.5*((x-v)/bandwidth)**2) for v in sampled) for x in xs]
        max_y=max(max_y,max(ys) if ys else 0); curves.append((label,ys))
    max_y=max_y or 1; width,height,left,right,top,bottom=900,430,74,25,62,62; plot_w=width-left-right; plot_h=height-top-bottom
    parts=["<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 900 430' role='img'>","<rect width='900' height='430' fill='white'/>","<text x='74' y='34' font-family='Microsoft YaHei,Arial' font-size='20' font-weight='700' fill='#173b31'>{}</text>".format(html.escape(title))]
    for i in range(5):
        y=top+plot_h*i/4
        parts.append("<line x1='{0}' y1='{1:.1f}' x2='{2}' y2='{1:.1f}' stroke='#e5ece8'/><text x='{3}' y='{4:.1f}' text-anchor='end' font-size='10' fill='#6b7e76'>{5:.3g}</text>".format(left,y,width-right,left-8,y+4,max_y*(1-i/4)))
    parts.append("<line x1='{0}' y1='{1}' x2='{2}' y2='{1}' stroke='#8ba097'/><line x1='{0}' y1='{3}' x2='{0}' y2='{1}' stroke='#8ba097'/>".format(left,height-bottom,width-right,top))
    for i in range(5):
        x=left+plot_w*i/4; value=lo+(hi-lo)*i/4
        parts.append("<text x='{:.1f}' y='390' text-anchor='middle' font-size='11' fill='#566b62'>{:.4g}</text>".format(x,value))
    for index,(label,ys) in enumerate(curves):
        points=" ".join("{:.1f},{:.1f}".format(left+plot_w*i/119,top+plot_h*(1-y/max_y)) for i,y in enumerate(ys)); color=palette[index%len(palette)]
        parts.append("<polyline points='{}' fill='none' stroke='{}' stroke-width='2.5' opacity='.9'/>".format(points,color))
        lx=650+(index%2)*115; ly=28+(index//2)*16
        parts.append("<line x1='{0}' y1='{1}' x2='{2}' y2='{1}' stroke='{3}' stroke-width='3'/><text x='{4}' y='{5}' font-size='10' fill='#40584f'>{6}</text>".format(lx,ly,lx+16,color,lx+20,ly+4,html.escape(str(label)[:16])))
    parts.append("<text x='470' y='418' text-anchor='middle' font-size='12' fill='#50665d'>{}</text><text transform='translate(18 230) rotate(-90)' text-anchor='middle' font-size='12' fill='#50665d'>密度</text></svg>".format(html.escape(unit or "表型值")))
    return "".join(parts)


def _threshold_catalog(species_id, custom):
    rows=[]
    for code,info in TRAIT_INFO.items(): rows.append({"trait":code,"name":info[0],"unit":info[1],**_threshold_for(code,species_id,custom)})
    for code in custom:
        if code not in TRAIT_INFO: rows.append({"trait":code,"name":code,"unit":"",**_threshold_for(code,species_id,custom)})
    return rows


class PhenotypeAnalyzer:
    def __init__(self):
        self._runs = {}

    @staticmethod
    def default_report_root():
        base=os.environ.get("LOCALAPPDATA") or str(Path.home()/".local"/"share")
        return Path(base)/"GPA-Accelerator"/"phenotype-reports"

    def catalog(self):
        data=phenotype_catalog(); data["default_report_root"]=str(self.default_report_root()); return data

    @staticmethod
    def _observation(row,mapping,sample,trait,value):
        def get(role): return row.get(mapping[role]) if mapping.get(role) else None
        return {"sample":sample,"trait":str(trait).strip(),"value":value,"year":str(get("year") or "").strip(),"location":str(get("location") or "").strip(),"latitude":_number(get("latitude")),"longitude":_number(get("longitude")),"replicate":str(get("replicate") or "").strip(),"group":str(get("group") or "").strip()}

    @staticmethod
    def _warning(level,trait,code,message,evidence,advice):
        return {"level":level,"trait":trait,"code":code,"message":message,"evidence":evidence,"advice":advice}

    def analyze(self,payload):
        path=Path(os.path.expandvars(os.path.expanduser(str(payload.get("path") or "")))).resolve()
        if not path.is_file(): raise VCFError("表型文件不存在：{}".format(path))
        if path.stat().st_size>1024*1024*1024: raise VCFError("表型文件超过1 GB；请拆分后分析")
        headers,raw_rows,source=_read_table(path,payload.get("sheet"))
        if not headers or not raw_rows: raise VCFError("表型文件没有可读取的数据行")
        mapping_config=payload.get("columns") if isinstance(payload.get("columns"),dict) else {}
        mapping={role:_detect_column(headers,role,mapping_config.get(role)) for role in ALIASES}
        table_format=str(payload.get("format") or "auto")
        if table_format=="auto": table_format="long" if mapping["trait"] and mapping["value"] else "wide"
        if not mapping["sample"]: raise VCFError("无法识别材料/样本列；请填写列映射")
        observations=[]; invalid_values=0; missing_values=0; attempted_values=0; metadata_columns={x for x in mapping.values() if x}
        requested=[x.strip() for x in str(payload.get("traits") or "").replace(",","\n").splitlines() if x.strip()]
        if table_format=="long":
            if not mapping["trait"] or not mapping["value"]: raise VCFError("长表需要性状列和数值列")
            for row in raw_rows:
                trait=str(row.get(mapping["trait"]) or "").strip(); sample=str(row.get(mapping["sample"]) or "").strip(); value=_number(row.get(mapping["value"]))
                if not trait or not sample: continue
                if requested and trait not in requested: continue
                attempted_values+=1
                if value is None:
                    if str(row.get(mapping["value"]) or "").strip().lower() in MISSING: missing_values+=1
                    else: invalid_values+=1
                    continue
                observations.append(self._observation(row,mapping,sample,trait,value))
        elif table_format=="wide":
            candidates=[x for x in headers if x not in metadata_columns]; trait_columns=[x for x in candidates if x in requested] if requested else []
            if not trait_columns:
                for column in candidates:
                    raw=[row.get(column) for row in raw_rows[:300] if str(row.get(column) or "").strip().lower() not in MISSING]
                    if raw and sum(_number(x) is not None for x in raw)/len(raw)>=.7: trait_columns.append(column)
            if not trait_columns: raise VCFError("宽表中没有识别出数值表型列")
            for row in raw_rows:
                sample=str(row.get(mapping["sample"]) or "").strip()
                if not sample: continue
                for trait in trait_columns:
                    attempted_values+=1; value=_number(row.get(trait))
                    if value is None:
                        if str(row.get(trait) or "").strip().lower() in MISSING: missing_values+=1
                        else: invalid_values+=1
                        continue
                    observations.append(self._observation(row,mapping,sample,trait,value))
        else: raise VCFError("表格格式只能是auto、wide或long")
        if not observations: raise VCFError("没有解析到有效数值表型")
        custom=_parse_custom_thresholds(payload.get("custom_thresholds")); species_id=str(payload.get("cotton_species") or "")
        if species_id and species_id not in COTTON_SPECIES: raise VCFError("未知棉花物种预设")
        config=payload.get("outlier") if isinstance(payload.get("outlier"),dict) else {}; shift_sd=max(.1,float(payload.get("shift_sd") or 1.0)); trend_sd=max(.01,float(payload.get("trend_sd_per_year") or .25))
        by_trait=defaultdict(list)
        for row in observations: by_trait[row["trait"]].append(row)
        root_text=str(payload.get("output_dir") or "").strip(); root=Path(os.path.expandvars(os.path.expanduser(root_text))).resolve() if root_text else self.default_report_root().resolve()
        run_id=time.strftime("%Y%m%d-%H%M%S-")+uuid.uuid4().hex[:8]; run_dir=root/run_id; run_dir.mkdir(parents=True,exist_ok=False)
        trait_results=[]; all_outliers=[]; estimates=[]; warnings=[]; aggregate_rows=[]; density_files=[]; density_mode=str(payload.get("density_mode") or "overlay")
        missing_rate=missing_values/max(1,attempted_values); invalid_rate=invalid_values/max(1,attempted_values)
        if missing_rate>.10: warnings.append(self._warning("warning","cohort","HIGH_PHENOTYPE_MISSING","表型缺失比例较高","{:.2%}的预期表型单元为空".format(missing_rate),"按材料、年份、地点和表型分层查看缺失机制；非随机缺失不可简单均值填补"))
        if invalid_rate>.01: warnings.append(self._warning("warning","cohort","NON_NUMERIC_VALUES","存在较多非数值内容","{:.2%}的预期表型单元无法解析为数值".format(invalid_rate),"检查单位字符、小数点、合并单元格和缺失值编码"))
        duplicate_keys=[(x["sample"],x["trait"],x.get("year"),x.get("location"),x.get("replicate")) for x in observations]
        duplicate_count=len(duplicate_keys)-len(set(duplicate_keys))
        if duplicate_count: warnings.append(self._warning("warning","cohort","DUPLICATE_OBSERVATIONS","存在重复观测键","{}条记录具有相同材料、性状、年份、地点和重复编号".format(duplicate_count),"核对是否为真实技术重复；不要在未确认前静默去重"))
        for trait in sorted(by_trait):
            rows=by_trait[trait]; values=[x["value"] for x in rows]; summary=_describe(values); threshold=_threshold_for(trait,species_id,custom)
            info=TRAIT_INFO.get(trait.upper())
            if info and info[1]=="%" and threshold.get("single") and summary.get("q95",0)<=1.5 and threshold["single"][0]>1.5:
                threshold["guide_single_percent"]=list(threshold["single"]); threshold["guide_mean_percent"]=list(threshold["mean"]) if threshold.get("mean") else None
                threshold["single"]=[x/100 for x in threshold["single"]]
                if threshold.get("mean"): threshold["mean"]=[x/100 for x in threshold["mean"]]
                threshold["detected_scale"]="fraction_0_to_1"
                warnings.append(self._warning("info",trait,"PERCENT_SCALE_INFERRED","检测到比例可能以0–1存储","数据P95≤1.5，已将指南百分数范围仅在本次比较中换算为0–1","确认原始单位；导出结果保留自动换算记录"))
            outliers,bounds=_outlier_rows(rows,config,threshold); all_outliers.extend(outliers); entries,variance=_blue_blup(rows)
            estimates.extend({"trait":trait,**item} for item in entries); years=_summary_groups(rows,"year"); locations=_summary_groups(rows,"location")
            for field,groups in (("year",years),("location",locations)):
                for item in groups: aggregate_rows.append({"trait":trait,"dimension":field,"level":item[field],**{k:v for k,v in item.items() if k!=field}})
            sample_env=defaultdict(list)
            for row in rows: sample_env[(row["sample"],row.get("year") or "",row.get("location") or "")].append(row["value"])
            replicates=[{"sample":k[0],"year":k[1],"location":k[2],"trait":trait,"n":len(v),"mean":statistics.fmean(v),"median":statistics.median(v),"sd":_safe_sd(v)} for k,v in sorted(sample_env.items())]
            overall_sd=summary.get("sd") or 0; yrange=max((x["mean"] for x in years),default=0)-min((x["mean"] for x in years),default=0); lrange=max((x["mean"] for x in locations),default=0)-min((x["mean"] for x in locations),default=0)
            year_shift=yrange/overall_sd if overall_sd and len(years)>1 else 0; location_shift=lrange/overall_sd if overall_sd and len(locations)>1 else 0
            slope=_linear_trend([(x["year"],x["mean"]) for x in years]); slope_z=slope/overall_sd if slope is not None and overall_sd else None
            site_slopes=[]
            for site in sorted({x.get("location") for x in rows if x.get("location")}):
                site_rows=[x for x in rows if x.get("location")==site]; site_years=_summary_groups(site_rows,"year"); site_slope=_linear_trend([(x["year"],x["mean"]) for x in site_years])
                if site_slope is not None and overall_sd: site_slopes.append((site,site_slope/overall_sd))
            if year_shift>shift_sd: warnings.append(self._warning("warning",trait,"YEAR_SHIFT","不同年份均值偏移较大","年度均值极差为总体SD的{:.2f}倍".format(year_shift),"检查年份环境、测量口径、批次与G×E；不要直接删除整年数据"))
            if location_shift>shift_sd: warnings.append(self._warning("warning",trait,"LOCATION_SHIFT","不同地点均值偏移较大","地点均值极差为总体SD的{:.2f}倍".format(location_shift),"检查地点环境、试验设计与空间效应"))
            if slope_z is not None and abs(slope_z)>trend_sd: warnings.append(self._warning("warning",trait,"TIME_TREND","表型存在明显年度趋势","标准化趋势为{:+.3f} SD/年".format(slope_z),"核对材料组成是否逐年变化；确认年份可比性"))
            if len(site_slopes)>1 and max(x[1] for x in site_slopes)-min(x[1] for x in site_slopes)>max(.5,2*trend_sd): warnings.append(self._warning("warning",trait,"SITE_TREND_DIVERGENCE","不同地点的年度趋势方向或幅度差异较大","地点标准化斜率范围{:+.3f}至{:+.3f} SD/年".format(min(x[1] for x in site_slopes),max(x[1] for x in site_slopes)),"检查年份×地点交互、材料不平衡和地点测量批次；建议使用多环境混合模型复核"))
            mean_range=threshold.get("mean")
            if mean_range and not mean_range[0]<=summary["mean"]<=mean_range[1]: warnings.append(self._warning("warning",trait,"BIOLOGICAL_MEAN_RANGE","总体均值超出指南参考范围","均值{:.4g}，参考{}–{}；置信度{}".format(summary["mean"],mean_range[0],mean_range[1],threshold["confidence"]),"先核对单位、性状定义和棉种；低置信范围仅作提示"))
            if outliers: warnings.append(self._warning("warning",trait,"OUTLIERS","检测到统计或生物学范围异常值","{}条被标记，其中{}条达到统计共识".format(len(outliers),sum(x["consensus_statistical"] for x in outliers)),"结合原始记录、重复和环境复核；不会自动删除"))
            if summary.get("sd")==0: warnings.append(self._warning("warning",trait,"ZERO_VARIANCE","表型没有可见变异","所有有效值相同","检查是否导入了编码列、舍入结果或错误的性状列"))
            if summary.get("n",0)<20: warnings.append(self._warning("info",trait,"SMALL_TRAIT_SAMPLE","表型有效观测较少","仅{}个有效观测".format(summary["n"]),"分布、正态性与离群判定不稳定，应结合原始记录解释"))
            unit=info[1] if info else ""
            if density_mode=="per_year" and years:
                grouped=defaultdict(list)
                for row in rows: grouped[str(row.get("year") or "未标年份")].append(row["value"])
                for year,vals in sorted(grouped.items()):
                    name="density_{}_{}.svg".format(_slug(trait),_slug(year)); (run_dir/name).write_text(_svg_density([(year,vals)],"{} · {} 密度分布".format(trait,year),unit),encoding="utf-8"); density_files.append({"trait":trait,"label":str(year),"name":name})
            else:
                if density_mode=="overlay" and years:
                    grouped=defaultdict(list)
                    for row in rows: grouped[str(row.get("year") or "未标年份")].append(row["value"])
                    series=sorted(grouped.items()); mode="overlay"
                else: series=[("全部",values)]; mode="pooled"
                name="density_{}_{}.svg".format(_slug(trait),mode); (run_dir/name).write_text(_svg_density(series,"{} 密度分布".format(trait),unit),encoding="utf-8"); density_files.append({"trait":trait,"label":mode,"name":name})
            geo=[{k:row.get(k) for k in ("sample","year","location","latitude","longitude","value")} for row in rows if row.get("latitude") is not None and row.get("longitude") is not None]
            trait_results.append({"trait":trait,"name":info[0] if info else trait,"unit":unit,"summary":summary,"threshold":threshold,"outlier_bounds":bounds,"outlier_count":len(outliers),"consensus_outlier_count":sum(x["consensus_statistical"] for x in outliers),"year_summaries":years,"location_summaries":locations,"replicate_summaries":replicates[:5000],"variance_components":variance,"time_trend":{"slope":slope,"standardized_slope_per_year":slope_z,"year_shift_sd":year_shift,"location_shift_sd":location_shift,"site_slopes":site_slopes},"geographic_points":geo[:5000]})
        correlations=_trait_correlations(observations)
        if len(by_trait)>1: (run_dir/"trait_correlation_heatmap.svg").write_text(_svg_correlations(correlations),encoding="utf-8")
        species=COTTON_SPECIES.get(species_id)
        result={"schema_version":"gpa-phenotype-qc-1.0","run_id":run_id,"generated_at":time.strftime("%Y-%m-%d %H:%M:%S"),"input":{"path":str(path),"name":path.name,"size":path.stat().st_size,"rows":len(raw_rows),"headers":headers,"source":source,"table_format":table_format,"mapping":mapping,"attempted_values":attempted_values,"missing_values":missing_values,"missing_rate":missing_rate,"invalid_values":invalid_values,"invalid_rate":invalid_rate,"duplicate_observation_keys":duplicate_count},"species":{"id":species_id or None,**({k:v for k,v in species.items() if k!="overrides"} if species else {"name":"通用/用户自定义"})},"summary":{"observations":len(observations),"samples":len({x["sample"] for x in observations}),"traits":len(by_trait),"years":len({x["year"] for x in observations if x.get("year")}),"locations":len({x["location"] for x in observations if x.get("location")}),"outliers":len(all_outliers),"warnings":len(warnings),"status":"warning" if any(x["level"]=="warning" for x in warnings) else "pass"},"configuration":{"density_mode":density_mode,"outlier":config,"shift_sd":shift_sd,"trend_sd_per_year":trend_sd,"custom_thresholds":custom,"automatic_deletion":False},"traits":trait_results,"correlations":correlations,"outliers":all_outliers,"warnings":warnings,"threshold_catalog":_threshold_catalog(species_id,custom),"method_notes":["IQR、3-sigma、MAD robust-z、尾部分位数和生物学范围并行计算。","离群值只是复核候选，GPA-Accelerator不会自动删除记录。","BLUE为环境中心化材料均值；BLUP为随机材料效应经验收缩估计。","数据不足或定义不清的表型不启用伪精确生物学范围。"]}
        artifacts=self._write_artifacts(run_dir,result,aggregate_rows,estimates,density_files); self._runs[run_id]={"dir":run_dir,"artifacts":{x["name"] for x in artifacts}}
        return {"run_id":run_id,"result":result,"artifacts":artifacts,"run_dir":str(run_dir)}

    @staticmethod
    def _write_tsv(path,rows,fields):
        with path.open("w",encoding="utf-8-sig",newline="") as handle:
            writer=csv.DictWriter(handle,fieldnames=fields,delimiter="\t",extrasaction="ignore"); writer.writeheader(); writer.writerows(rows)

    def _write_artifacts(self,run_dir,result,aggregate_rows,estimates,density_files):
        summary_path=run_dir/"phenotype_summary.json"; trait_path=run_dir/"trait_statistics.tsv"; aggregate_path=run_dir/"environment_summaries.tsv"
        replicate_path=run_dir/"replicate_averages.tsv"; estimate_path=run_dir/"blue_blup_estimates.tsv"; outlier_path=run_dir/"outlier_candidates.tsv"
        warning_path=run_dir/"phenotype_warnings.tsv"; threshold_path=run_dir/"cotton_thresholds.tsv"; correlation_path=run_dir/"trait_correlations.tsv"; report_path=run_dir/"phenotype_report.html"
        summary_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
        trait_rows=[]; replicate_rows=[]
        for item in result["traits"]:
            trait_rows.append({"trait":item["trait"],"name":item["name"],"unit":item["unit"],**item["summary"],**item["variance_components"],**item["time_trend"],"outlier_count":item["outlier_count"],"consensus_outlier_count":item["consensus_outlier_count"]})
            replicate_rows.extend(item["replicate_summaries"])
        trait_fields=["trait","name","unit","n","mean","median","trimmed_mean_5pct","sd","variance","se","cv","mad","iqr","min","q05","q25","q75","q95","max","skewness","excess_kurtosis","genetic_variance","residual_variance","broad_sense_h2_entry_mean","harmonic_replicates","environments","slope","standardized_slope_per_year","year_shift_sd","location_shift_sd","outlier_count","consensus_outlier_count"]
        self._write_tsv(trait_path,trait_rows,trait_fields)
        self._write_tsv(aggregate_path,aggregate_rows,["trait","dimension","level","n","mean","median","trimmed_mean_5pct","sd","variance","se","cv","mad","iqr","min","q05","q25","q75","q95","max","skewness","excess_kurtosis"])
        self._write_tsv(replicate_path,replicate_rows,["sample","year","location","trait","n","mean","median","sd"])
        self._write_tsv(estimate_path,estimates,["trait","sample","n","raw_mean","blue","blup","blup_reliability"])
        self._write_tsv(outlier_path,result["outliers"],["row","sample","trait","value","year","location","methods","method_count","consensus_statistical","severity"])
        self._write_tsv(warning_path,result["warnings"],["level","trait","code","message","evidence","advice"])
        self._write_tsv(correlation_path,result.get("correlations",[]),["trait_x","trait_y","n","pearson_r","spearman_rho"])
        threshold_rows=[]
        for item in result["threshold_catalog"]:
            single=item.get("single") or [None,None]; mean=item.get("mean") or [None,None]
            threshold_rows.append({**item,"single_min":single[0],"single_max":single[1],"mean_min":mean[0],"mean_max":mean[1]})
        self._write_tsv(threshold_path,threshold_rows,["trait","name","unit","single_min","single_max","mean_min","mean_max","confidence","source"])
        report_path.write_text(self._render_report(result,density_files,run_dir),encoding="utf-8")
        core=[report_path,summary_path,trait_path,aggregate_path,replicate_path,estimate_path,outlier_path,warning_path,threshold_path,correlation_path]
        figures=sorted(run_dir.glob("*.svg")); zip_path=run_dir/"GPA_Accelerator_phenotype_QC.zip"
        with zipfile.ZipFile(zip_path,"w",compression=zipfile.ZIP_DEFLATED) as archive:
            for item in core+figures: archive.write(item,item.name)
        artifacts=[]
        for item in core+figures+[zip_path]:
            base="/api/phenotype/artifact?run_id={}&name={}".format(result["run_id"],item.name)
            artifacts.append({"name":item.name,"size":item.stat().st_size,"url":base,"view_url":base+"&view=1"})
        return artifacts

    def artifact(self,run_id,name):
        run=self._runs.get(str(run_id or ""))
        if not run: raise VCFError("表型分析任务不存在或本地服务已重启")
        if name not in run["artifacts"] or Path(name).name!=name: raise VCFError("表型报告文件不存在")
        path=(run["dir"]/name).resolve()
        if path.parent!=run["dir"].resolve() or not path.is_file(): raise VCFError("表型报告文件不存在")
        return path

    @staticmethod
    def _render_report(result,density_files,run_dir):
        trait_rows="".join("<tr><td>{}</td><td>{}</td><td>{}</td><td>{:.4g}</td><td>{:.4g}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(html.escape(x["trait"]),html.escape(x["name"]),x["summary"]["n"],x["summary"]["mean"],x["summary"]["sd"],"{:.3f}".format(x["variance_components"]["broad_sense_h2_entry_mean"]) if x["variance_components"]["broad_sense_h2_entry_mean"] is not None else "—",x["outlier_count"],html.escape(x["threshold"]["confidence"])) for x in result["traits"])
        warning_rows="".join("<tr><td><span class='badge'>{}</span></td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(html.escape(x["level"]),html.escape(x["trait"]),html.escape(x["message"]),html.escape(x["evidence"]),html.escape(x["advice"])) for x in result["warnings"]) or "<tr><td colspan='5'>未发现达到当前规则的报警。</td></tr>"
        density_options=[]
        for item in density_files:
            encoded=base64.b64encode((run_dir/item["name"]).read_bytes()).decode("ascii")
            density_options.append("<option value='data:image/svg+xml;base64,{}'>{} · {}</option>".format(encoded,html.escape(item["trait"]),html.escape(item["label"])))
        density_options="".join(density_options)
        correlation_path=run_dir/"trait_correlation_heatmap.svg"
        correlation_src="data:image/svg+xml;base64,"+base64.b64encode(correlation_path.read_bytes()).decode("ascii") if correlation_path.is_file() else ""
        payload=json.dumps({"traits":result["traits"]},ensure_ascii=False).replace("</","<\\/")
        template="""<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>GPA-Accelerator 表型质量报告</title><style>
body{margin:0;background:#eef4f0;color:#17372e;font-family:system-ui,'Microsoft YaHei',sans-serif}main{max-width:1180px;margin:auto;padding:28px}.hero,.card{margin-bottom:18px;padding:24px;border:1px solid #d6e2dc;border-radius:18px;background:white}.hero{color:white;background:linear-gradient(135deg,#153f33,#387e69)}h1,h2{margin-top:0}.kpis{display:grid;grid-template-columns:repeat(6,1fr);gap:10px}.kpi{padding:13px;border-radius:12px;background:#edf6f1}.kpi b{display:block;font-size:23px}.grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}.table{overflow:auto;max-height:560px}table{width:100%;border-collapse:collapse;font-size:12px}th,td{padding:8px;border-bottom:1px solid #e5ece8;text-align:left;white-space:nowrap}.badge{padding:3px 7px;border-radius:10px;background:#fff0c9;color:#805300}select,input,button{min-height:38px;padding:7px 10px;border:1px solid #bfcfc7;border-radius:8px;background:white}button{cursor:pointer}.chart{min-height:340px;border:1px solid #dde7e2;border-radius:12px;background:#fbfdfc}.chart img{width:100%}.controls{display:flex;gap:8px;align-items:center;flex-wrap:wrap;margin-bottom:12px}.controls label{display:grid;gap:3px;font-size:11px}.muted{color:#687b73;font-size:12px}.geo svg,.chart svg{width:100%;height:340px}.point{stroke:white;stroke-width:1;opacity:.82}@media(max-width:800px){.kpis{grid-template-columns:repeat(2,1fr)}.grid{grid-template-columns:1fr}}@media print{body{background:white}main{max-width:none;padding:0}.controls{display:none}.card,.hero{break-inside:avoid}}
</style></head><body><main><section class='hero'><p>GPA-Accelerator · PHENOTYPE QUALITY REPORT</p><h1>表型质量评估、BLUE/BLUP与多环境诊断</h1><p>__INPUT__ · __GENERATED__</p></section><section class='card'><div class='kpis'><div class='kpi'>观测值<b>__OBS__</b></div><div class='kpi'>材料<b>__SAMPLES__</b></div><div class='kpi'>表型<b>__TRAITS__</b></div><div class='kpi'>年份<b>__YEARS__</b></div><div class='kpi'>地点<b>__LOCATIONS__</b></div><div class='kpi'>离群候选<b>__OUTLIERS__</b></div></div><p class='muted'>物种预设：__SPECIES__。离群值仅作复核候选，本报告不会自动删除原始数据。</p></section>
<section class='card'><h2>密度分布图</h2><div class='controls'><label>选择表型/年份<select id='densitySelect'>__DENSITY_OPTIONS__</select></label><button onclick='window.print()'>打印/另存为PDF</button></div><div class='chart'><img id='densityImage' alt='表型密度分布'></div></section><section class='card'><h2>表型相关性</h2><p class='muted'>基于每个材料跨环境均值进行Pearson与Spearman计算；热图显示Pearson相关。</p><div class='chart' id='correlationChart'><img src='__CORR_SRC__' alt='表型相关热图'></div></section>
<section class='grid'><div class='card'><h2>时间动态</h2><div class='controls'><label>表型<select id='traitSelect'></select></label><label>年份<input id='yearSlider' type='range' min='0' value='0'></label><button id='playBtn'>播放/暂停</button><b id='yearLabel'></b></div><div class='chart' id='timeChart'></div></div><div class='card'><h2>地理分布动态</h2><p class='muted'>需要经纬度列；点颜色与大小表示当前年份表型值。</p><div class='chart' id='geoChart'></div></div></section>
<section class='card'><h2>表型统计、遗传力与离群值</h2><div class='table'><table><thead><tr><th>表型</th><th>名称</th><th>N</th><th>均值</th><th>SD</th><th>H²</th><th>离群</th><th>阈值置信度</th></tr></thead><tbody>__TRAIT_ROWS__</tbody></table></div></section>
<section class='card'><h2>报警与复核建议</h2><div class='table'><table><thead><tr><th>级别</th><th>表型</th><th>问题</th><th>证据</th><th>建议</th></tr></thead><tbody>__WARNING_ROWS__</tbody></table></div></section>
<section class='card'><h2>下载说明</h2><p>完整ZIP包含统计、原始重复均值、BLUE/BLUP、离群候选、阈值表、报警和全部SVG图；请从GPA-Accelerator主页面下载。</p></section></main><script>
const DATA=__PAYLOAD__;const density=document.getElementById('densitySelect'),img=document.getElementById('densityImage');function setDensity(){img.src=density.value}density.onchange=setDensity;if(density.options.length)setDensity();
const traits=DATA.traits||[],ts=document.getElementById('traitSelect'),slider=document.getElementById('yearSlider'),label=document.getElementById('yearLabel');traits.forEach((x,i)=>ts.add(new Option(x.trait,i)));function esc(x){return String(x??'').replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]))}
function render(){const t=traits[+ts.value||0];if(!t)return;const ys=t.year_summaries||[],idx=Math.min(+slider.value,Math.max(0,ys.length-1));slider.max=Math.max(0,ys.length-1);label.textContent=ys[idx]?.year||'全部';const w=520,h=300,p=48;let vals=ys.map(x=>x.mean);if(!vals.length)vals=[t.summary.mean];let lo=Math.min(...vals),hi=Math.max(...vals);if(lo===hi){lo-=.5;hi+=.5}const pts=vals.map((v,i)=>`${p+(w-2*p)*i/Math.max(1,vals.length-1)},${h-p-(h-2*p)*(v-lo)/(hi-lo)}`).join(' ');document.getElementById('timeChart').innerHTML=`<svg viewBox="0 0 ${w} ${h}"><line x1="${p}" y1="${h-p}" x2="${w-p}" y2="${h-p}" stroke="#9aaba3"/><line x1="${p}" y1="${p}" x2="${p}" y2="${h-p}" stroke="#9aaba3"/><polyline points="${pts}" fill="none" stroke="#176b5a" stroke-width="3"/>${vals.map((v,i)=>`<circle cx="${p+(w-2*p)*i/Math.max(1,vals.length-1)}" cy="${h-p-(h-2*p)*(v-lo)/(hi-lo)}" r="5" fill="${i===idx?'#d36b2c':'#176b5a'}"><title>${esc(ys[i]?.year||'全部')}: ${v.toFixed(4)}</title></circle>`).join('')}</svg>`;let geo=(t.geographic_points||[]).filter(x=>!ys.length||String(x.year)===String(ys[idx]?.year));if(!geo.length){document.getElementById('geoChart').innerHTML='<p style="padding:30px" class="muted">当前表型/年份没有经纬度数据。</p>';return}const lons=geo.map(x=>x.longitude),lats=geo.map(x=>x.latitude),gvals=geo.map(x=>x.value),xmin=Math.min(...lons),xmax=Math.max(...lons),ymin=Math.min(...lats),ymax=Math.max(...lats),vmin=Math.min(...gvals),vmax=Math.max(...gvals);document.getElementById('geoChart').innerHTML=`<svg viewBox="0 0 520 340"><rect x="35" y="25" width="450" height="270" fill="#eef5f1" stroke="#c8d8d0"/>${geo.map(x=>{const cx=50+420*(x.longitude-xmin)/Math.max(.0001,xmax-xmin),cy=280-240*(x.latitude-ymin)/Math.max(.0001,ymax-ymin),q=(x.value-vmin)/Math.max(.0001,vmax-vmin),r=5+7*q,c=`rgb(${Math.round(42+190*q)},${Math.round(132-55*q)},${Math.round(102-55*q)})`;return `<circle class="point" cx="${cx}" cy="${cy}" r="${r}" fill="${c}"><title>${esc(x.location)} · ${esc(x.sample)}: ${x.value}</title></circle>`}).join('')}</svg>`}
ts.onchange=()=>{slider.value=0;render()};slider.oninput=render;let timer;document.getElementById('playBtn').onclick=()=>{if(timer){clearInterval(timer);timer=null;return}timer=setInterval(()=>{slider.value=(+slider.value+1)%(+slider.max+1);render()},900)};render();
</script></body></html>"""
        values={"__INPUT__":html.escape(result["input"]["name"]),"__GENERATED__":html.escape(result["generated_at"]),"__OBS__":"{:,}".format(result["summary"]["observations"]),"__SAMPLES__":"{:,}".format(result["summary"]["samples"]),"__TRAITS__":str(result["summary"]["traits"]),"__YEARS__":str(result["summary"]["years"]),"__LOCATIONS__":str(result["summary"]["locations"]),"__OUTLIERS__":str(result["summary"]["outliers"]),"__SPECIES__":html.escape(result["species"].get("name") or "通用/用户自定义"),"__DENSITY_OPTIONS__":density_options,"__CORR_SRC__":correlation_src,"__TRAIT_ROWS__":trait_rows,"__WARNING_ROWS__":warning_rows,"__PAYLOAD__":payload}
        for key,value in values.items(): template=template.replace(key,value)
        if not correlation_src: template=template.replace("<img src='' alt='表型相关热图'>","<p class='muted' style='padding:30px'>至少需要两个表型才能绘制相关热图。</p>")
        return template
