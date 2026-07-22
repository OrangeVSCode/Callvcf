#!/usr/bin/env python3
"""Single-locus genotype-to-phenotype association for GPA-Accelerator."""

import base64
import csv
import html
import json
import math
import re
import statistics
import time
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path

from vcf_service import VCFError, parse_loci


MISSING = {"", ".", "na", "nan", "null", "none", "n/a", "-", "--"}
SAMPLE_ALIASES = {
    "sample", "sampleid", "sample_id", "id", "iid", "material", "materialid",
    "material_id", "accession", "line", "genotype", "品种", "材料", "材料名", "样本",
}
TRAIT_ALIASES = {"trait", "phenotype", "variable", "性状", "表型", "指标"}
VALUE_ALIASES = {"value", "phenotypevalue", "phenotype_value", "measurement", "数值", "值", "表型值"}
YEAR_ALIASES = {"year", "season", "年份", "年度"}
LOCATION_ALIASES = {"location", "site", "environment", "env", "地点", "试点", "环境"}


def _norm(value):
    return re.sub(r"[\s_\-./()（）]+", "", str(value or "").strip().casefold())


def _number(value):
    text = str(value if value is not None else "").strip()
    if text.casefold() in MISSING:
        return None
    try:
        value = float(text.replace(",", "").rstrip("%"))
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def _safe_sd(values):
    return statistics.stdev(values) if len(values) > 1 else 0.0


def _describe(values):
    if not values:
        return {"n": 0, "mean": None, "median": None, "sd": None, "se": None, "min": None, "max": None}
    sd = _safe_sd(values)
    return {
        "n": len(values), "mean": statistics.fmean(values), "median": statistics.median(values),
        "sd": sd, "se": sd / math.sqrt(len(values)) if values else None,
        "min": min(values), "max": max(values),
    }


def _betacf(a, b, x):
    maximum, epsilon, floor = 200, 3e-14, 1e-300
    qab, qap, qam = a + b, a + 1.0, a - 1.0
    c = 1.0
    d = 1.0 - qab * x / qap
    d = 1.0 / (floor if abs(d) < floor else d)
    h = d
    for index in range(1, maximum + 1):
        m2 = 2 * index
        aa = index * (b - index) * x / ((qam + m2) * (a + m2))
        d = 1.0 + aa * d; d = floor if abs(d) < floor else d
        c = 1.0 + aa / c; c = floor if abs(c) < floor else c
        d = 1.0 / d; h *= d * c
        aa = -(a + index) * (qab + index) * x / ((a + m2) * (qap + m2))
        d = 1.0 + aa * d; d = floor if abs(d) < floor else d
        c = 1.0 + aa / c; c = floor if abs(c) < floor else c
        d = 1.0 / d
        delta = d * c; h *= delta
        if abs(delta - 1.0) < epsilon:
            break
    return h


def _regularized_beta(x, a, b):
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    front = math.exp(math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log1p(-x))
    if x < (a + 1.0) / (a + b + 2.0):
        return front * _betacf(a, b, x) / a
    return 1.0 - front * _betacf(b, a, 1.0 - x) / b


def _student_p(t_value, degrees):
    if degrees <= 0 or t_value is None:
        return None
    x = degrees / (degrees + t_value * t_value)
    return max(0.0, min(1.0, _regularized_beta(x, degrees / 2.0, 0.5)))


def _f_survival(f_value, df1, df2):
    if f_value is None or f_value < 0 or df1 <= 0 or df2 <= 0:
        return None
    x = df2 / (df2 + df1 * f_value)
    return max(0.0, min(1.0, _regularized_beta(x, df2 / 2.0, df1 / 2.0)))


def _gamma_q(a, x):
    if x < 0 or a <= 0:
        return None
    if x == 0:
        return 1.0
    epsilon, maximum, floor = 3e-14, 200, 1e-300
    if x < a + 1.0:
        term = total = 1.0 / a
        ap = a
        for _ in range(maximum):
            ap += 1.0; term *= x / ap; total += term
            if abs(term) < abs(total) * epsilon:
                break
        p_value = total * math.exp(-x + a * math.log(x) - math.lgamma(a))
        return max(0.0, min(1.0, 1.0 - p_value))
    b = x + 1.0 - a; c = 1.0 / floor; d = 1.0 / b; h = d
    for index in range(1, maximum + 1):
        an = -index * (index - a); b += 2.0
        d = an * d + b; d = floor if abs(d) < floor else d
        c = b + an / c; c = floor if abs(c) < floor else c
        d = 1.0 / d; delta = d * c; h *= delta
        if abs(delta - 1.0) < epsilon:
            break
    return max(0.0, min(1.0, math.exp(-x + a * math.log(x) - math.lgamma(a)) * h))


def _linear_regression(pairs):
    n = len(pairs)
    if n < 3:
        return {"n": n, "slope": None, "se": None, "pvalue": None, "r2": None, "standardized_beta": None}
    xs = [float(x) for x, _ in pairs]; ys = [float(y) for _, y in pairs]
    mean_x, mean_y = statistics.fmean(xs), statistics.fmean(ys)
    sxx = sum((x - mean_x) ** 2 for x in xs); syy = sum((y - mean_y) ** 2 for y in ys)
    sxy = sum((x - mean_x) * (y - mean_y) for x, y in pairs)
    if sxx <= 0 or syy <= 0:
        return {"n": n, "slope": None, "se": None, "pvalue": None, "r2": None, "standardized_beta": None}
    slope = sxy / sxx; intercept = mean_y - slope * mean_x
    sse = sum((y - intercept - slope * x) ** 2 for x, y in pairs)
    se = math.sqrt(max(0.0, sse / max(1, n - 2) / sxx))
    t_value = slope / se if se > 0 else None
    return {
        "n": n, "slope": slope, "se": se, "ci95_low": slope - 1.96 * se,
        "ci95_high": slope + 1.96 * se, "pvalue": _student_p(t_value, n - 2),
        "r2": max(0.0, min(1.0, sxy * sxy / (sxx * syy))),
        "standardized_beta": slope * math.sqrt(sxx / syy),
    }


def _anova(groups):
    groups = [values for values in groups if values]
    total_n, count = sum(map(len, groups)), len(groups)
    if count < 2 or total_n <= count:
        return {"f": None, "df_between": None, "df_within": None, "pvalue": None, "eta_squared": None}
    grand = statistics.fmean([value for group in groups for value in group])
    between = sum(len(group) * (statistics.fmean(group) - grand) ** 2 for group in groups)
    within = sum(sum((value - statistics.fmean(group)) ** 2 for value in group) for group in groups)
    df1, df2 = count - 1, total_n - count
    f_value = (between / df1) / (within / df2) if within > 0 else None
    return {
        "f": f_value, "df_between": df1, "df_within": df2,
        "pvalue": _f_survival(f_value, df1, df2),
        "eta_squared": between / (between + within) if between + within > 0 else None,
    }


def _kruskal(groups):
    indexed = [(value, group_index) for group_index, group in enumerate(groups) for value in group]
    if len(indexed) < 3 or sum(bool(group) for group in groups) < 2:
        return {"h": None, "df": None, "pvalue": None}
    indexed.sort(key=lambda item: item[0]); rank_sums = [0.0] * len(groups); tie_sum = 0.0
    cursor = 0
    while cursor < len(indexed):
        end = cursor + 1
        while end < len(indexed) and indexed[end][0] == indexed[cursor][0]:
            end += 1
        rank = (cursor + 1 + end) / 2.0
        for _, group_index in indexed[cursor:end]:
            rank_sums[group_index] += rank
        size = end - cursor; tie_sum += size ** 3 - size; cursor = end
    n = len(indexed)
    h_value = 12.0 / (n * (n + 1.0)) * sum(rank_sums[i] ** 2 / len(groups[i]) for i in range(len(groups)) if groups[i]) - 3.0 * (n + 1.0)
    correction = 1.0 - tie_sum / (n ** 3 - n) if n > 1 else 1.0
    h_value = h_value / correction if correction > 0 else h_value
    degrees = sum(bool(group) for group in groups) - 1
    return {"h": h_value, "df": degrees, "pvalue": _gamma_q(degrees / 2.0, h_value / 2.0)}


def _bh_adjust(items, p_key="pvalue"):
    valid = sorted([(index, item.get(p_key)) for index, item in enumerate(items) if item.get(p_key) is not None], key=lambda item: item[1])
    adjusted = [None] * len(items); running = 1.0; total = len(valid)
    for rank in range(total, 0, -1):
        index, pvalue = valid[rank - 1]; running = min(running, pvalue * total / rank); adjusted[index] = min(1.0, running)
    for item, qvalue in zip(items, adjusted):
        item["fdr_qvalue"] = qvalue


def _decode_text(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise VCFError("无法识别表型文本编码；建议另存为 UTF-8 CSV/TSV/TXT")


def _read_table(path, sheet=None):
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise VCFError("读取 XLSX 需要 openpyxl；Windows 完整版已内置")
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_name = sheet if sheet in workbook.sheetnames else workbook.sheetnames[0]
        iterator = workbook[sheet_name].iter_rows(values_only=True)
        try:
            headers = [str(value or "").strip() for value in next(iterator)]
        except StopIteration:
            workbook.close(); raise VCFError("表型文件为空")
        rows = [dict(zip(headers, values)) for values in iterator if any(value is not None and str(value).strip() for value in values)]
        sheets = list(workbook.sheetnames); workbook.close()
        return headers, rows, {"format": "xlsx", "sheet": sheet_name, "sheets": sheets, "header_inferred": False}
    if suffix == ".xls":
        raise VCFError("旧版 .XLS 请另存为 .XLSX、CSV 或 TSV")
    if suffix not in {".csv", ".tsv", ".txt", ".ps", ".phen", ".pheno"}:
        raise VCFError("表型关联支持 .xlsx、.csv、.tsv、.txt、.ps、.phen 和 .pheno")
    text, encoding = _decode_text(path)
    lines = [line for line in text.splitlines() if line.strip() and not line.lstrip().startswith("#")]
    if not lines:
        raise VCFError("表型文件为空")
    delimiter = None
    try:
        delimiter = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=",\t;").delimiter
    except csv.Error:
        pass
    parsed = list(csv.reader(lines, delimiter=delimiter)) if delimiter else [re.split(r"\s+", line.strip()) for line in lines]
    parsed = [[str(value).strip() for value in row] for row in parsed]
    width = max(map(len, parsed)); parsed = [row + [""] * (width - len(row)) for row in parsed]
    first, second = parsed[0], parsed[1] if len(parsed) > 1 else []
    known = {_norm(value) for value in first} & {_norm(value) for value in SAMPLE_ALIASES | TRAIT_ALIASES | VALUE_ALIASES | YEAR_ALIASES | LOCATION_ALIASES}
    looks_header = bool(known) or (bool(second) and any(_number(first[i]) is None and _number(second[i]) is not None for i in range(1, min(len(first), len(second)))))
    if looks_header:
        headers, data = first, parsed[1:]
        inferred = False
    else:
        if width == 2:
            headers = ["ID", path.stem]
        elif width == 3:
            headers = ["FID", "ID", path.stem]
        else:
            raise VCFError("无表头文本超过3列；请添加表头并明确材料列与表型列")
        data = parsed; inferred = True
    rows = [dict(zip(headers, row)) for row in data if any(str(value).strip() for value in row)]
    return headers, rows, {"format": suffix.lstrip("."), "encoding": encoding, "delimiter": delimiter or "whitespace", "header_inferred": inferred}


def _resolve_column(headers, requested, aliases, required=False, label="列"):
    by_fold = {header.casefold(): header for header in headers}
    if requested:
        resolved = by_fold.get(str(requested).strip().casefold())
        if not resolved:
            raise VCFError("未找到{}：{}".format(label, requested))
        return resolved
    normalized = {_norm(header): header for header in headers}
    for alias in aliases:
        if _norm(alias) in normalized:
            return normalized[_norm(alias)]
    if required:
        raise VCFError("无法自动识别{}；请手动填写列名".format(label))
    return None


def _is_replicate(name):
    return bool(re.search(r"_(?:R|REP)\d+$", str(name), re.I))


def _is_average(name):
    return bool(re.search(r"_(?:AV|AVG|MEAN)$", str(name), re.I))


def _trait_base(name):
    return re.sub(r"_(?:AV|AVG|MEAN|R\d+|REP\d+)$", "", str(name).strip(), flags=re.I)


def _resolve_traits(available, query, mode, include_replicates=False):
    query = str(query or "").strip()
    if not available:
        raise VCFError("表型文件中没有可分析的数值表型")
    if not query and len(available) == 1:
        return list(available)
    if not query:
        raise VCFError("检测到多个表型；请输入精确表型名或前缀")
    if mode == "exact":
        direct = [name for name in available if name.casefold() == query.casefold()]
        if direct:
            return direct
        base = [name for name in available if _trait_base(name).casefold() == query.casefold()]
        averages = [name for name in base if _is_average(name)]
        if averages:
            return averages
        non_repeats = [name for name in base if not _is_replicate(name)]
        if non_repeats:
            return non_repeats
        if include_replicates and base:
            return base
        raise VCFError("未找到精确表型“{}”；可改用前缀模式".format(query))
    matches = [name for name in available if str(name).casefold().startswith(query.casefold())]
    if not matches:
        raise VCFError("没有以“{}”开头的表型".format(query))
    if include_replicates:
        return matches
    averages = [name for name in matches if _is_average(name)]
    if averages:
        return averages
    representatives = [name for name in matches if not _is_replicate(name)]
    return representatives or matches


def _genotype(gt):
    raw = str(gt or ".").split(":", 1)[0].replace("|", "/")
    alleles = raw.split("/")
    if not alleles or any(allele == "." or not allele.isdigit() for allele in alleles):
        return None, "MISSING", raw, None
    integers = [int(allele) for allele in alleles]
    dosage = sum(allele > 0 for allele in integers); ploidy = len(integers)
    if dosage == 0:
        category = "HOM_REF"
    elif dosage == ploidy:
        category = "HOM_ALT"
    else:
        category = "HET"
    return dosage, category, raw, ploidy


def _composites(traits, trait_values, min_coverage):
    if len(traits) < 2:
        return {}
    samples = sorted({sample for trait in traits for sample in trait_values[trait]})
    minimum = max(1, math.ceil(len(traits) * min_coverage))
    raw = {}
    means = {trait: statistics.fmean(trait_values[trait].values()) for trait in traits if trait_values[trait]}
    sds = {trait: _safe_sd(list(trait_values[trait].values())) for trait in traits if trait_values[trait]}
    standardized = {}
    for sample in samples:
        values = [trait_values[trait][sample] for trait in traits if sample in trait_values[trait]]
        if len(values) >= minimum:
            raw[sample] = statistics.fmean(values)
        z_values = [(trait_values[trait][sample] - means[trait]) / sds[trait] for trait in traits if sample in trait_values[trait] and sds.get(trait)]
        if len(z_values) >= minimum:
            standardized[sample] = statistics.fmean(z_values)
    return {"COMBINED_RAW_MEAN": raw, "COMBINED_ENV_Z": standardized}


def _analyze_trait(label, kind, values, genotype_map, min_group_size):
    groups = defaultdict(list); samples_by_group = defaultdict(list); pairs = []; sample_rows = []
    ploidies = []
    for vcf_sample, phenotype_value in values.items():
        dosage, category, gt, ploidy = _genotype(genotype_map.get(vcf_sample))
        if dosage is None:
            continue
        ploidies.append(ploidy); groups[category].append(phenotype_value); samples_by_group[category].append(vcf_sample)
        pairs.append((dosage, phenotype_value))
        sample_rows.append({"sample": vcf_sample, "trait": label, "analysis_kind": kind, "gt": gt, "genotype_group": category, "alt_dosage": dosage, "phenotype": phenotype_value})
    order = ["HOM_REF", "HET", "HOM_ALT"]
    summaries = {group: _describe(groups[group]) for group in order}
    regression = _linear_regression(pairs)
    anova = _anova([groups[group] for group in order])
    kruskal = _kruskal([groups[group] for group in order])
    valid_counts = [len(groups[group]) for group in order if groups[group]]
    power_warning = len(valid_counts) < 2 or min(valid_counts, default=0) < min_group_size or len(pairs) < max(10, min_group_size * 2)
    ref_mean, alt_mean = summaries["HOM_REF"]["mean"], summaries["HOM_ALT"]["mean"]
    return {
        "trait": label, "analysis_kind": kind, "n": len(pairs), "genotype_groups": summaries,
        "sample_examples": {group: samples_by_group[group][:30] for group in order},
        "dosage_regression": regression, "anova": anova, "kruskal_wallis": kruskal,
        "homozygote_difference_alt_minus_ref": alt_mean - ref_mean if alt_mean is not None and ref_mean is not None else None,
        "ploidy_values": sorted(set(ploidies)), "low_power_warning": power_warning,
        "sample_rows": sample_rows,
    }


def _svg_plot(variant_key, analyses):
    shown = analyses[:16]; width = 1120; row_height = 68; height = 100 + row_height * len(shown)
    all_values = [group["mean"] for item in shown for group in item["genotype_groups"].values() if group["mean"] is not None]
    if not all_values:
        return '<svg xmlns="http://www.w3.org/2000/svg" width="900" height="180"><text x="30" y="80">没有可绘制的表型值</text></svg>'
    low, high = min(all_values), max(all_values)
    if low == high:
        low -= 0.5; high += 0.5
    left, right = 280, width - 70
    colors = {"HOM_REF": "#276f61", "HET": "#d1913c", "HOM_ALT": "#b94b5f"}
    labels = {"HOM_REF": "0/0", "HET": "0/1", "HOM_ALT": "1/1"}
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="#fbfdfb"/>', f'<text x="28" y="34" font-size="22" font-weight="700" fill="#173c34">{html.escape(variant_key)}：基因型组表型均值 ± 95% CI</text>']
    for tick in range(6):
        value = low + (high - low) * tick / 5; x = left + (right - left) * tick / 5
        parts += [f'<line x1="{x:.1f}" y1="55" x2="{x:.1f}" y2="{height-30}" stroke="#dfe8e3"/>', f'<text x="{x:.1f}" y="{height-10}" text-anchor="middle" font-size="12" fill="#526962">{value:.3g}</text>']
    for index, item in enumerate(shown):
        center = 78 + index * row_height
        parts.append(f'<text x="270" y="{center+5}" text-anchor="end" font-size="14" fill="#253f38">{html.escape(item["trait"][:34])}</text>')
        for offset, group_name in zip((-17, 0, 17), ("HOM_REF", "HET", "HOM_ALT")):
            summary = item["genotype_groups"][group_name]
            if summary["mean"] is None:
                continue
            mean, se = summary["mean"], summary["se"] or 0.0; y = center + offset
            x = left + (mean - low) / (high - low) * (right - left)
            x1 = left + (mean - 1.96 * se - low) / (high - low) * (right - left)
            x2 = left + (mean + 1.96 * se - low) / (high - low) * (right - left)
            x1, x2 = max(left, x1), min(right, x2)
            parts += [f'<line x1="{x1:.1f}" y1="{y}" x2="{x2:.1f}" y2="{y}" stroke="{colors[group_name]}" stroke-width="3"/>', f'<circle cx="{x:.1f}" cy="{y}" r="5" fill="{colors[group_name]}"><title>{labels[group_name]} n={summary["n"]}, mean={mean:.5g}</title></circle>']
    parts.append('</svg>')
    return "".join(parts)


class VariantPhenotypeAnalyzer:
    def __init__(self, service):
        self.service = service
        self._runs = {}

    def analyze(self, payload):
        vcf_path = payload.get("vcf_path") or payload.get("path")
        phenotype_path = Path(str(payload.get("phenotype_path") or "")).expanduser().resolve()
        if not phenotype_path.is_file():
            raise VCFError("表型文件不存在：{}".format(phenotype_path))
        loci = parse_loci(payload.get("locus"), limit=20)
        if len(loci) != 1:
            raise VCFError("当前界面一次分析一个位点；批量位点将在 EMMAX/全基因组模块中处理")
        mode = str(payload.get("match_mode") or "exact").strip().casefold()
        if mode not in {"exact", "prefix"}:
            raise VCFError("表型匹配方式必须是 exact 或 prefix")
        query = str(payload.get("trait_query") or "").strip()
        include_replicates = bool(payload.get("include_replicates"))
        combine = bool(payload.get("combine", mode == "prefix"))
        min_coverage = max(0.1, min(1.0, float(payload.get("min_trait_coverage", 0.5))))
        min_group_size = max(2, int(payload.get("min_group_size", 5)))
        headers, rows, source = _read_table(phenotype_path, str(payload.get("sheet") or "").strip() or None)
        columns = payload.get("columns") or {}
        sample_column = _resolve_column(headers, columns.get("sample"), SAMPLE_ALIASES, True, "材料/样本列")
        trait_column = _resolve_column(headers, columns.get("trait"), TRAIT_ALIASES)
        value_column = _resolve_column(headers, columns.get("value"), VALUE_ALIASES)
        year_column = _resolve_column(headers, columns.get("year"), YEAR_ALIASES)
        location_column = _resolve_column(headers, columns.get("location"), LOCATION_ALIASES)
        year_filter = str(payload.get("year_filter") or "").strip().casefold()
        location_filter = str(payload.get("location_filter") or "").strip().casefold()
        metadata_columns = {column for column in (sample_column, trait_column, value_column, year_column, location_column) if column}
        values_by_trait = defaultdict(lambda: defaultdict(list))
        if trait_column and value_column:
            available = sorted({str(row.get(trait_column) or "").strip() for row in rows if str(row.get(trait_column) or "").strip()})
            selected_traits = _resolve_traits(available, query, mode, include_replicates)
            selected_fold = {trait.casefold(): trait for trait in selected_traits}
            for row in rows:
                sample = str(row.get(sample_column) or "").strip(); trait = str(row.get(trait_column) or "").strip()
                if not sample or trait.casefold() not in selected_fold:
                    continue
                if year_filter and str(row.get(year_column) or "").strip().casefold() != year_filter:
                    continue
                if location_filter and str(row.get(location_column) or "").strip().casefold() != location_filter:
                    continue
                value = _number(row.get(value_column))
                if value is not None:
                    values_by_trait[selected_fold[trait.casefold()]][sample].append(value)
        else:
            numeric_candidates = [header for header in headers if header not in metadata_columns and any(_number(row.get(header)) is not None for row in rows)]
            selected_traits = _resolve_traits(numeric_candidates, query, mode, include_replicates)
            for row in rows:
                sample = str(row.get(sample_column) or "").strip()
                if not sample:
                    continue
                if year_filter and year_column and str(row.get(year_column) or "").strip().casefold() != year_filter:
                    continue
                if location_filter and location_column and str(row.get(location_column) or "").strip().casefold() != location_filter:
                    continue
                for trait in selected_traits:
                    value = _number(row.get(trait))
                    if value is not None:
                        values_by_trait[trait][sample].append(value)
        trait_values = {trait: {sample: statistics.fmean(values) for sample, values in sample_rows.items()} for trait, sample_rows in values_by_trait.items() if sample_rows}
        selected_traits = [trait for trait in selected_traits if trait in trait_values]
        if not selected_traits:
            raise VCFError("所选表型没有有效数值；请检查 NA、筛选条件和列映射")
        metadata, _, vcf_samples, records = self.service.query_records(vcf_path, payload.get("locus"))
        if not records:
            raise VCFError("VCF 中不存在位点 {}:{}".format(*loci[0]))
        phenotype_names = {sample.casefold(): sample for trait in trait_values.values() for sample in trait}
        matched = {vcf_sample: phenotype_names[vcf_sample.casefold()] for vcf_sample in vcf_samples if vcf_sample.casefold() in phenotype_names}
        if not matched:
            raise VCFError("VCF 样本名与表型材料名没有交集；请检查 ID 列和命名规则")
        aligned = {trait: {vcf_sample: trait_values[trait][phenotype_sample] for vcf_sample, phenotype_sample in matched.items() if phenotype_sample in trait_values[trait]} for trait in selected_traits}
        analysis_values = [(trait, "separate", aligned[trait]) for trait in selected_traits]
        if combine:
            for name, values in _composites(selected_traits, aligned, min_coverage).items():
                analysis_values.append(("{}__{}".format(query or "TRAITS", name), "combined_raw" if name.endswith("RAW_MEAN") else "combined_z", values))
        variants = []; all_samples = []
        for record in records:
            analyses = [_analyze_trait(label, kind, values, record["genotypes"], min_group_size) for label, kind, values in analysis_values]
            _bh_adjust([item["dosage_regression"] for item in analyses])
            for item in analyses:
                item["fdr_qvalue"] = item["dosage_regression"].get("fdr_qvalue")
                all_samples.extend({"variant": record["key"], **row} for row in item.pop("sample_rows"))
            variants.append({
                "record": {key: record.get(key) for key in ("key", "chrom", "pos", "id", "ref", "alt", "variant_type", "svtype", "svlen")},
                "analyses": analyses,
            })
        phenotype_samples = set(phenotype_names)
        matched_fraction = len(matched) / max(1, min(len(vcf_samples), len(phenotype_samples)))
        warnings = []
        if matched_fraction < 0.8:
            warnings.append("VCF与表型材料匹配率低于80%；请核对材料命名")
        if any("," in str(variant["record"].get("alt") or "") for variant in variants):
            warnings.append("位点含多个ALT；剂量按任意非REF等位基因拷贝数计算")
        if combine and len(selected_traits) > 1:
            warnings.append("原值综合均值只适合同量纲表型；跨环境比较优先使用按环境标准化的 COMBINED_ENV_Z")
        if any(item["low_power_warning"] for variant in variants for item in variant["analyses"]):
            warnings.append("部分分析的基因型组样本量较小；P值与效应方向应谨慎解释")
        best = min((item for variant in variants for item in variant["analyses"] if item["dosage_regression"].get("pvalue") is not None), key=lambda item: item["dosage_regression"]["pvalue"], default=None)
        result = {
            "schema_version": "gpa-variant-phenotype-1.0", "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "input": {"vcf": metadata["path"], "phenotype": str(phenotype_path), "phenotype_source": source, "sample_column": sample_column, "trait_column": trait_column, "value_column": value_column, "year_column": year_column, "location_column": location_column},
            "query": {"locus": "{}:{}".format(*loci[0]), "trait_query": query, "match_mode": mode, "include_replicates": include_replicates, "combine": combine, "min_trait_coverage": min_coverage, "year_filter": year_filter, "location_filter": location_filter},
            "summary": {"vcf_samples": len(vcf_samples), "phenotype_samples": len(phenotype_samples), "matched_samples": len(matched), "matched_fraction": matched_fraction, "selected_traits": len(selected_traits), "variant_records": len(records), "analyses": sum(len(variant["analyses"]) for variant in variants)},
            "selected_traits": selected_traits, "variants": variants, "warnings": warnings,
            "best_association": {"trait": best["trait"], "pvalue": best["dosage_regression"]["pvalue"], "fdr_qvalue": best.get("fdr_qvalue"), "slope": best["dosage_regression"]["slope"]} if best else None,
            "method_notes": ["每个材料在同一表型中的重复值先取平均，避免伪重复。", "剂量回归检验每增加一个ALT等位基因的表型变化；同时给出单因素ANOVA和Kruskal-Wallis。", "多个环境分别分析时对剂量回归P值进行Benjamini-Hochberg FDR校正。", "前缀综合同时给出原值材料均值与按环境z标准化后的综合指数。", "这是候选位点的描述/验证分析，不替代含亲缘矩阵和群体结构协变量的全基因组混合模型。"],
        }
        run_id = time.strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:8]
        output_root = Path(str(payload.get("output_dir") or (Path.home() / "GPA-Accelerator" / "association-reports"))).expanduser().resolve()
        run_dir = output_root / run_id; run_dir.mkdir(parents=True, exist_ok=False)
        result["run_id"] = run_id
        artifacts = self._write_artifacts(run_dir, result, all_samples)
        self._runs[run_id] = {"dir": run_dir, "artifacts": {item["name"] for item in artifacts}}
        return {"run_id": run_id, "run_dir": str(run_dir), "result": result, "artifacts": artifacts}

    @staticmethod
    def _write_tsv(path, rows, fields):
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", extrasaction="ignore")
            writer.writeheader(); writer.writerows(rows)

    def _write_artifacts(self, run_dir, result, sample_rows):
        json_path = run_dir / "association_summary.json"
        json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        association_rows = []
        group_rows = []
        for variant in result["variants"]:
            for item in variant["analyses"]:
                regression, anova, kw = item["dosage_regression"], item["anova"], item["kruskal_wallis"]
                association_rows.append({"variant": variant["record"]["key"], "trait": item["trait"], "analysis_kind": item["analysis_kind"], "n": item["n"], "alt_dosage_slope": regression["slope"], "slope_se": regression["se"], "pvalue": regression["pvalue"], "fdr_qvalue": item["fdr_qvalue"], "r2": regression["r2"], "standardized_beta": regression["standardized_beta"], "anova_f": anova["f"], "anova_pvalue": anova["pvalue"], "eta_squared": anova["eta_squared"], "kruskal_h": kw["h"], "kruskal_pvalue": kw["pvalue"], "homozygote_difference_alt_minus_ref": item["homozygote_difference_alt_minus_ref"], "low_power_warning": item["low_power_warning"]})
                for group, summary in item["genotype_groups"].items():
                    group_rows.append({"variant": variant["record"]["key"], "trait": item["trait"], "analysis_kind": item["analysis_kind"], "genotype_group": group, **summary})
        self._write_tsv(run_dir / "trait_associations.tsv", association_rows, list(association_rows[0]) if association_rows else ["variant", "trait"])
        self._write_tsv(run_dir / "genotype_group_statistics.tsv", group_rows, list(group_rows[0]) if group_rows else ["variant", "trait"])
        self._write_tsv(run_dir / "sample_genotype_phenotypes.tsv", sample_rows, ["variant", "sample", "trait", "analysis_kind", "gt", "genotype_group", "alt_dosage", "phenotype"])
        first_variant = result["variants"][0]
        svg = _svg_plot(first_variant["record"]["key"], first_variant["analyses"])
        (run_dir / "genotype_phenotype_plot.svg").write_text(svg, encoding="utf-8")
        report = self._report_html(result, svg, association_rows)
        (run_dir / "variant_phenotype_report.html").write_text(report, encoding="utf-8")
        zip_path = run_dir / "GPA_Accelerator_variant_phenotype_association.zip"
        names = ["association_summary.json", "trait_associations.tsv", "genotype_group_statistics.tsv", "sample_genotype_phenotypes.tsv", "genotype_phenotype_plot.svg", "variant_phenotype_report.html"]
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name in names:
                archive.write(run_dir / name, name)
        names.append(zip_path.name)
        artifacts = []
        for name in names:
            path = run_dir / name; base = "/api/variant-phenotype/artifact?run_id={}&name={}".format(result["run_id"], name)
            artifacts.append({"name": name, "size": path.stat().st_size, "url": base, "view_url": base + "&view=1"})
        return artifacts

    @staticmethod
    def _report_html(result, svg, association_rows):
        rows = "".join("<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(html.escape(str(row["trait"])), html.escape(str(row["analysis_kind"])), row["n"], "—" if row["alt_dosage_slope"] is None else "{:.5g}".format(row["alt_dosage_slope"]), "—" if row["pvalue"] is None else "{:.4g}".format(row["pvalue"]), "—" if row["fdr_qvalue"] is None else "{:.4g}".format(row["fdr_qvalue"]), "—" if row["eta_squared"] is None else "{:.4g}".format(row["eta_squared"])) for row in association_rows)
        warnings = "".join("<li>{}</li>".format(html.escape(item)) for item in result["warnings"]) or "<li>未触发额外警告</li>"
        encoded = base64.b64encode(svg.encode("utf-8")).decode("ascii")
        report = """<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width\"><title>GPA-Accelerator 位点×表型报告</title><style>body{font-family:Arial,'Microsoft YaHei',sans-serif;margin:0;background:#eef4f0;color:#173c34}.page{max-width:1220px;margin:auto;padding:30px}.card{background:white;border:1px solid #d7e3dc;border-radius:18px;padding:24px;margin:18px 0;overflow:auto}h1,h2{margin-top:0}table{border-collapse:collapse;width:100%}th,td{padding:10px;border-bottom:1px solid #e1e9e4;text-align:left}th{background:#edf5f1}img{max-width:100%;height:auto}.meta{display:flex;gap:18px;flex-wrap:wrap}.meta b{font-size:24px;display:block}</style></head><body><div class=\"page\"><h1>位点 × 表型关联报告</h1><p>__LOCUS__ · __TRAIT__ · __GENERATED__</p><div class=\"card meta\"><span><b>__MATCHED__</b>匹配材料</span><span><b>__TRAITS__</b>分开表型</span><span><b>__ANALYSES__</b>统计分析</span></div><div class=\"card\"><h2>基因型分组图</h2><img src=\"data:image/svg+xml;base64,__SVG__\"></div><div class=\"card\"><h2>关联统计</h2><table><thead><tr><th>表型</th><th>类型</th><th>N</th><th>ALT剂量效应</th><th>P</th><th>FDR</th><th>η²</th></tr></thead><tbody>__ROWS__</tbody></table></div><div class=\"card\"><h2>解释与警告</h2><ul>__WARNINGS__</ul><ul>__NOTES__</ul></div></div></body></html>"""
        replacements = {
            "__LOCUS__": html.escape(result["query"]["locus"]), "__TRAIT__": html.escape(result["query"]["trait_query"]),
            "__GENERATED__": html.escape(result["generated_at"]), "__MATCHED__": str(result["summary"]["matched_samples"]),
            "__TRAITS__": str(result["summary"]["selected_traits"]), "__ANALYSES__": str(result["summary"]["analyses"]),
            "__SVG__": encoded, "__ROWS__": rows, "__WARNINGS__": warnings,
            "__NOTES__": "".join("<li>{}</li>".format(html.escape(item)) for item in result["method_notes"]),
        }
        for token, value in replacements.items():
            report = report.replace(token, value)
        return report

    def artifact(self, run_id, name):
        run = self._runs.get(str(run_id or ""))
        if not run:
            raise VCFError("位点×表型分析任务不存在或本地服务已重启")
        if name not in run["artifacts"] or Path(name).name != name:
            raise VCFError("位点×表型报告文件不存在")
        path = (run["dir"] / name).resolve()
        if run["dir"] not in path.parents or not path.is_file():
            raise VCFError("禁止访问分析目录外的文件")
        return path
